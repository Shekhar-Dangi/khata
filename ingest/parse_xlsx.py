"""HDFC & Indian Bank .xlsx statement parsers -> canonical transaction rows.

Both are tabular (unlike the Slice PDF), but differ:
  HDFC        header: Date | Narration | Chq./Ref.No. | Value Dt | Withdrawal Amt. | Deposit Amt. | Closing Balance
              dates 'DD/MM/YY' strings; separate withdrawal/deposit columns; numeric balance.
  Indian Bank header: Txn Date | Description | Cheque No | Debit Amount | Credit Amount | Balance
              dates are datetime cells; separate debit/credit columns; balance like '962.380 CR'.

The bank is detected from the header signature. Canonical output per txn:
txn_date (YYYY-MM-DD), amount_paise (signed: +credit / -debit), narration, ref,
bank_balance_paise, type='regular' — same shape the Slice parser and import path use.

Usage: python parse_xlsx.py <path-to-xlsx>
Local-only.
"""

import re
import sys
from datetime import datetime
from decimal import Decimal

import openpyxl

HDFC_HEADER = ["Date", "Narration", "Chq./Ref.No.", "Withdrawal Amt.", "Deposit Amt.", "Closing Balance"]
INDIAN_HEADER = ["Txn Date", "Description", "Debit Amount", "Credit Amount", "Balance"]


def to_paise(value):
    """Rupees -> signed int paise. Accepts a number, or a string like '962.380 CR'
    / '1,234.56' / '100.00 DR' (DR -> negative). None/blank -> None."""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return int(round(Decimal(str(value)) * 100))
    s = str(value).strip()
    if not s:
        return None
    sign = -1 if s.upper().endswith("DR") else 1
    cleaned = re.sub(r"[^\d.]", "", s)  # strip ₹, commas, CR/DR, spaces
    if not cleaned:
        return None
    return sign * int(round(Decimal(cleaned) * 100))


def find_header_row(ws, header):
    """1-based index of the row containing every column name in `header`, else None."""
    want = [h.lower() for h in header]
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        cells = [str(c).strip().lower() for c in row if c is not None]
        if all(h in cells for h in want):
            return i
    return None


def _signed(credit_paise, debit_paise):
    if credit_paise is not None:
        return credit_paise
    if debit_paise is not None:
        return -debit_paise
    return None


def parse_hdfc(ws):
    start = find_header_row(ws, HDFC_HEADER) + 1
    txns = []
    for row in ws.iter_rows(min_row=start, values_only=True):
        cell = row[0]
        if cell is None or not re.match(r"^\d{2}/\d{2}/\d{2}$", str(cell).strip()):
            continue  # skip '****' separators, summary, footer — data rows lead with DD/MM/YY
        txns.append({
            "txn_date": datetime.strptime(str(cell).strip(), "%d/%m/%y").strftime("%Y-%m-%d"),
            "amount_paise": _signed(to_paise(row[5]), to_paise(row[4])),  # deposit / withdrawal
            "narration": row[1],
            "ref": row[2] or None,
            "bank_balance_paise": to_paise(row[6]),
            "type": "regular",
        })
    return txns


def parse_indian(ws):
    start = find_header_row(ws, INDIAN_HEADER) + 1
    txns = []
    for row in ws.iter_rows(min_row=start, values_only=True):
        cell = row[0]
        if not isinstance(cell, datetime):  # data rows carry a datetime; footer/blanks don't
            continue
        txns.append({
            "txn_date": cell.strftime("%Y-%m-%d"),
            "amount_paise": _signed(to_paise(row[4]), to_paise(row[3])),  # credit / debit
            "narration": row[1],
            "ref": row[2] or None,
            "bank_balance_paise": to_paise(row[5]),
            "type": "regular",
        })
    return txns


def parse(path):
    ws = openpyxl.load_workbook(path, data_only=True).active
    if find_header_row(ws, HDFC_HEADER):
        return "hdfc", parse_hdfc(ws)
    if find_header_row(ws, INDIAN_HEADER):
        return "indian_bank", parse_indian(ws)
    raise ValueError("unrecognized statement format (no known header row found)")


def main():
    bank, txns = parse(sys.argv[1])
    print(f"detected bank: {bank}")
    print(f"parsed {len(txns)} transactions\n")
    for t in txns:
        amt = t["amount_paise"] / 100 if t["amount_paise"] is not None else None
        bal = t["bank_balance_paise"] / 100 if t["bank_balance_paise"] is not None else None
        print(f"{t['txn_date']}  amt={amt:>12}  bal={bal:>12}  ref={t['ref']}")
        print(f"            {t['narration']}")


if __name__ == "__main__":
    main()
